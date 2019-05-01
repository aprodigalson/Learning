s = 'hello\n你好'
with open('test.txt','w') as fp:
    fp.write(s)

with open('test.txt') as fp:
    print(fp.read())

def demo9_2(src,dst,srcEncoding,dstEncoding):
    with open(src,'r',encoding=srcEncoding) as srcfp:
        with open(dst,'w',encoding=dstEncoding) as dstfp:
            dstfp.write(srcfp.read())

def demo9_3(filename):
    with open(filename,'r') as fp:
        for line in fp:
            print(line)

import json 
import csv
import fileinput
import linecache

with fileinput.input(files=('test_new.txt','test.txt')) as fp:
    for line in fp:
        print(line)

def demo9_13(fileName,infectedContent):
    with open(fileName,'r')as fp:
        lines = fp.readlines()
    for index,line in enumerate(lines):
        if line.strip().lower().startswith('<html>'):
            lines.insert(index+1,infectedContent)
            break
    with open(fileName,'w')as fp:
        fp.writelines(lines)

import chardet 

import pickle
i = 123
a = 66.6
s = '红框'
lst = [[12,2,3],[2.3,4]]
tu = (1,2,3)
coll = {2,3,4}
dic = {1:'appple',2:'bananaa'}
data = (i,a,s,lst,tu,coll,dic)

with open('sample_pickle.dat','wb') as f:
    try:
        pickle.dump(len(data),f)
        for item in data:
            pickle.dump(item,f)
    except:
        print('写入异常')

with open('sample_pickle.dat','rb') as f:
    n = pickle.load(f)
    for i in range(n):
        x = pickle.load(f)
        print(x)

import random 
import os
from docx import Document

columnsNumber = 4
def demo9_22(rowsNumber=20,grade=4):
    if grade < 3 :
        operators = '+-'
        biggest = 20
    elif grade <= 4:
        operators = '+-×÷'
        biggest = 100
    elif grade == 5:
        operators = '+-×÷('
        biggest = 100
    
    document = Document()
    table = document.add_table(rows=rowsNumber,cols=columnsNumber)
    for row in range(rowsNumber):
        for col in range(columnsNumber):
            first = random.randint(1,biggest)
            second = random.randint(1,biggest)
            operator = random.choice(operators)
            if operator !='(':
                if operator=='-':
                    if first<second:
                        first,second = second,first
                r = str(first).ljust(2,' ')+' '+operator+str(second).ljust(2,' ')+'='
            else:
                third = random.randint(1,100)
                while True:
                    o1 = random.choice(operators)
                    o2 = random.choice(operators)
                    if o1 !='(' and o2 != '(':
                        break
                rr = random.randint(1,100)
                if rr>50:
                    if o2=='-':
                        if second<third:
                            second,third = third,second
                    r = str(first).ljust(2,' ')+o1+'('+str(second).ljust(2,' ')+o2+str(third).ljust(2,' ')+')='
                else:
                    if o1=='-':
                        if first<second:
                            first,second=second,first        
                    r = '('+str(first).ljust(2,' ')+o1+str(second).ljust(2,' ')+')'+o2+str(third).ljust(2,' ')+'='
            cell=table.cell(row,col)
            cell.text = r
    document.save('kousuan.docx')
    os.startfile('kousuan.docx')

import zipfile
def demo9_24(filename):
    with zipfile.ZipFile(filename) as fp:
        for fn in fp.namelist():
            print(fn)

import sys
import time

pdfs = (pdfs for pdfs in os.listdir('.') if pdfs.endswith('.pdf'))

for pdf1 in pdfs:
    pdf = pdf1.replace(' ','_').replace('-','_').replace('&','_')
    os.rename(pdf1,pdf)
    print('='*30+'\n',pdf)
    txt = pdf[:-4]+'.txt'
    pdf2txt=os.path.dirname(sys.executable)
    pdf2txt=pdf2txt+'\\scripts\\pdf2txt.py'
    try:
        cmd = "python "+pdf2txt+' '+pdf+'>'+txt
        os.popen(cmd)
        time.sleep(2)
        with open(txt,encoding='utf-8')as fp:
            print(fp.read(20))
    except:
        pass


from string import ascii_letters
from os import listdir,rename
from os.path import join,splitext
from random import choice,randint

def demo10_1(directory):
    for fn in listdir(directory):
        _,ext = splitext(fn)
        n = randint(5,20)
        newName = ''.join((choice(ascii_letters) for _ in range(n)))
        rename(join(directory,fn),join(directory,newName+ext))


import sys
import zlib
import hashlib
import os.path

def demo10_2(filename):
    if os.path.isfile(filename):
        with open(filename,'rb') as fp:
            contexts = fp.read()
        return (zlib.crc32(contexts),hashlib.md5(contexts).hexdigest())
    else:
        print('file not exists')
    
import filecmp
import shutil
def demo10_3(srcDir,dstDir):
    if ((not os.path.isdir(srcDir)) or (not os.path.isdir(dstDir)) or (os.path.abspath(srcDir)!=srcDir) or (os.path.abspath(dstDir)!=dstDir)):
        usage()

    if os.path.commonpath((srcDir,dstDir))==srcDir:
        usage()
    for item in os.listdir(srcDir):
        srcItem = os.path.join(srcDir,item)
        dstItem = srcItem.replace(srcDir,dstDir)
        if os.path.isdir(srcItem):
            if not os.path.exists(dstItem):
                os.makedirs(dstItem)
                print('make directory' + dstItem)
            demo10_3(srcItem,dstItem)
        elif os.path.isfile(srcItem):
            if((not os.path.exists(dstItem)) or (not filecmp.cmp(srcItem,dstItem,shallow=False))):
                shutil.copyfile(srcItem,dstItem)
                print('file:'+srcItem+'==>'+dstItem)


def usage():
    print("error")
    sys.exit(0)

totalSize = 0
fileNum = 0
dirNum = 0
def demo10_4(path):
    global totalSize
    global fileNum
    global dirNum
    for lists in os.listdir(path):
        sub_path = os.path.join(path,lists)
        if os.path.isfile(sub_path):
            fileNum +=1
            totalSize += os.path.getsize(sub_path)
        elif os.path.isdir(sub_path):
            dirNum+=1
            demo10_4(sub_path)


